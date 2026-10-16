#!/usr/bin/env python3
"""Extract text for content checks by rendering a PDF and running Windows Scan text.

Searchable and scanned PDFs take the same path: each page becomes a PNG, then
Windows.Media.Ocr reads the picture. That is the engine behind Photos > Scan text.
The text already stored in the PDF is ignored. The result is a TXT file and a DOCX
file. Each PDF page's text stays on a Word page of the same size.

Usage:
    .\\.venv\\Scripts\\python.exe convert.py "D:\\path\\to\\file.pdf"
"""

from __future__ import annotations

import argparse
import asyncio
import sys
from pathlib import Path

import pymupdf

from scan_text import PlacedLine, ensure_models, recognize_page, scan_engine

# Windows Scan text rejects bitmaps whose longest side is above this.
_MAX_SIDE = 10000


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Render a PDF to images and extract text with Windows Scan text."
    )
    parser.add_argument("pdf", type=Path, help="Input PDF path")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=None,
        help="Output directory (default: output/<pdf name>)",
    )
    parser.add_argument(
        "--dpi",
        type=int,
        default=300,
        help="Render resolution (default: 300)",
    )
    return parser


def _page_dpi(page: pymupdf.Page, dpi: int) -> int:
    """Lower the render DPI when a page would exceed the Scan text size limit."""
    longest = max(page.rect.width, page.rect.height, 1) * dpi / 72
    if longest > _MAX_SIDE:
        return max(72, int(dpi * _MAX_SIDE / longest))
    return dpi


def render_pdf(pdf: Path, pages_dir: Path, dpi: int) -> list[Path]:
    """Write one PNG per page. The PDF text layer is ignored."""
    pages_dir.mkdir(parents=True, exist_ok=True)
    doc = pymupdf.open(pdf)
    paths: list[Path] = []
    try:
        for index, page in enumerate(doc, start=1):
            out = pages_dir / f"page_{index:04d}.png"
            pixmap = page.get_pixmap(dpi=_page_dpi(page, dpi), alpha=False)
            pixmap.save(out)
            paths.append(out)
            print(f"image {index}/{doc.page_count}: {out.name}", flush=True)
    finally:
        doc.close()
    return paths


def _pdf_page_sizes(pdf: Path) -> list[tuple[float, float]]:
    """Page width and height in points, one pair per PDF page."""
    doc = pymupdf.open(pdf)
    try:
        return [(float(page.rect.width), float(page.rect.height)) for page in doc]
    finally:
        doc.close()


def _line_units(line: str) -> float:
    """Approximate how wide a line is, in ems, so it can be fitted on the page."""
    units = 0.0
    for char in line:
        code = ord(char)
        if char == "\t":
            units += 2.0
        elif (
            0x1100 <= code <= 0x11FF
            or 0x2E80 <= code <= 0x9FFF
            or 0xAC00 <= code <= 0xD7AF
            or 0x3040 <= code <= 0x30FF
            or 0x0B80 <= code <= 0x0BFF
        ):
            units += 1.05
        elif char in "mwMW@#%&GOQD":
            units += 0.85
        elif char.isspace():
            units += 0.35
        else:
            units += 0.58
    return max(units, 1.0)


def _fit_page(lines: list[str], width: float, height: float) -> tuple[float, float, float]:
    """Return margin, font size, and exact line spacing that keep every line on one page."""
    margin = min(14.0, width * 0.04, height * 0.025)
    margin = max(6.0, margin)
    usable_w = max(20.0, width - (2 * margin) - 4.0)
    usable_h = max(20.0, height - (2 * margin) - 8.0)
    slot = usable_h / max(len(lines), 1)
    widest = max(_line_units(line) for line in lines)
    font = min(11.0, slot * 0.82, usable_w / widest)
    return margin, max(font, 1.0), slot


def _set_run_font(run, font_size: float) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.size = Pt(font_size)
    run.font.name = "Calibri"
    r_pr = run._element.get_or_add_rPr()
    fonts = r_pr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.append(fonts)
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:cs"), "Calibri")
    fonts.set(qn("w:eastAsia"), "Malgun Gothic")


def _set_line(paragraph, line: str, font_size: float, slot: float) -> None:
    from docx.enum.text import WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    paragraph.text = line
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(slot)
    fmt.left_indent = Pt(0)
    fmt.right_indent = Pt(0)
    fmt.first_line_indent = Pt(0)
    p_pr = paragraph._p.get_or_add_pPr()
    snap = p_pr.find(qn("w:snapToGrid"))
    if snap is None:
        snap = OxmlElement("w:snapToGrid")
        p_pr.append(snap)
    snap.set(qn("w:val"), "0")
    widow = p_pr.find(qn("w:widowControl"))
    if widow is None:
        widow = OxmlElement("w:widowControl")
        p_pr.append(widow)
    widow.set(qn("w:val"), "0")
    for run in paragraph.runs:
        _set_run_font(run, font_size)


def _png_size(path: Path) -> tuple[int, int]:
    import struct

    with path.open("rb") as handle:
        handle.seek(16)
        width, height = struct.unpack(">II", handle.read(8))
    return int(width), int(height)


def _frame_line(paragraph, line: PlacedLine, scale_x: float, scale_y: float, page_w: float, page_h: float, gap_pt: float) -> None:
    """Place one OCR line on the page at the same position it had in the PDF."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    x_pt = max(0.0, line.x * scale_x)
    y_pt = max(0.0, line.y * scale_y)
    h_pt = max(1.0, line.h * scale_y)
    units = _line_units(line.text)
    room = max(8.0, gap_pt)
    font = min(h_pt * 0.78, room / units, 12.0)
    font = max(font, 1.0)
    w_pt = min(max(line.w * scale_x, units * font) + 1.5, room)
    if y_pt + h_pt > page_h:
        y_pt = max(0.0, page_h - h_pt)
    if x_pt + w_pt > page_w:
        x_pt = max(0.0, page_w - w_pt)
    _set_line(paragraph, line.text, font, h_pt)
    p_pr = paragraph._p.get_or_add_pPr()
    frame = OxmlElement("w:framePr")
    frame.set(qn("w:w"), str(max(1, int(w_pt * 20))))
    frame.set(qn("w:h"), str(max(1, int(h_pt * 20))))
    frame.set(qn("w:hRule"), "exact")
    frame.set(qn("w:hAnchor"), "page")
    frame.set(qn("w:vAnchor"), "page")
    frame.set(qn("w:x"), str(max(0, int(x_pt * 20))))
    frame.set(qn("w:y"), str(max(0, int(y_pt * 20))))
    frame.set(qn("w:wrap"), "around")
    frame.set(qn("w:hSpace"), "0")
    frame.set(qn("w:vSpace"), "0")
    p_pr.append(frame)


def _row_gaps(lines: list[PlacedLine], scale_x: float, page_w: float) -> list[float]:
    """How much horizontal room each line has before the next line on its row."""
    rows: list[list[int]] = []
    order = sorted(range(len(lines)), key=lambda i: (lines[i].y, lines[i].x))
    for index in order:
        line = lines[index]
        if rows and line.y <= lines[rows[-1][-1]].y + lines[rows[-1][-1]].h * 0.55:
            rows[-1].append(index)
        else:
            rows.append([index])
    gaps = [page_w - (lines[i].x * scale_x) - 4.0 for i in range(len(lines))]
    for row in rows:
        row.sort(key=lambda i: lines[i].x)
        for pos, index in enumerate(row):
            x_pt = lines[index].x * scale_x
            if pos + 1 < len(row):
                gaps[index] = max(8.0, lines[row[pos + 1]].x * scale_x - x_pt - 3.0)
            else:
                gaps[index] = max(8.0, page_w - x_pt - 4.0)
    return gaps


def write_docx(
    path: Path,
    pages: list[tuple[int, str]],
    page_sizes: list[tuple[float, float]],
    layouts: list[list[PlacedLine]] | None = None,
    image_sizes: list[tuple[int, int]] | None = None,
) -> None:
    """Write one Word file. Each PDF page's text stays on one Word page of the same size."""
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    sizes = list(page_sizes) or [(612.0, 792.0)]
    if len(sizes) < len(pages):
        sizes.extend([sizes[-1]] * (len(pages) - len(sizes)))

    used_first = False
    for index, (_number, text) in enumerate(pages):
        width, height = sizes[index]
        lines = text.splitlines() or [""]
        margin, font_size, slot = _fit_page(lines, width, height)
        if index == 0:
            section = document.sections[0]
        elif (width, height) != sizes[index - 1]:
            section = document.add_section()
        else:
            section = document.sections[-1]
        section.page_width = Pt(width)
        section.page_height = Pt(height)
        section.left_margin = Pt(margin)
        section.right_margin = Pt(margin)
        section.top_margin = Pt(margin)
        section.bottom_margin = Pt(margin)
        section.header_distance = Pt(0)
        section.footer_distance = Pt(0)
        placed = layouts[index] if layouts and index < len(layouts) else []
        image_size = image_sizes[index] if image_sizes and index < len(image_sizes) else (0, 0)
        last = None
        if placed and image_size[0] and image_size[1]:
            section.left_margin = Pt(6)
            section.right_margin = Pt(6)
            section.top_margin = Pt(6)
            section.bottom_margin = Pt(6)
            scale_x = width / image_size[0]
            scale_y = height / image_size[1]
            gaps = _row_gaps(placed, scale_x, width)
            for line, gap in zip(placed, gaps):
                if not used_first and document.paragraphs:
                    last = document.paragraphs[0]
                else:
                    last = document.add_paragraph()
                used_first = True
                _frame_line(last, line, scale_x, scale_y, width, height, gap)
        else:
            for line in lines:
                if not used_first and document.paragraphs:
                    last = document.paragraphs[0]
                else:
                    last = document.add_paragraph()
                used_first = True
                _set_line(last, line, font_size, slot)
        next_size = sizes[index + 1] if index < len(pages) - 1 else None
        if last is not None and next_size == (width, height):
            breaker = document.add_paragraph()
            breaker.paragraph_format.space_before = Pt(0)
            breaker.paragraph_format.space_after = Pt(0)
            breaker.add_run().add_break(WD_BREAK.PAGE)
    document.save(path)


async def extract(pdf: Path, out_dir: Path, dpi: int) -> tuple[Path, Path]:
    pages_dir = out_dir / "pages"
    images = render_pdf(pdf, pages_dir, dpi)
    engine = await scan_engine()
    tessdata = ensure_models()
    parts: list[str] = []
    pages: list[tuple[int, str]] = []
    layouts: list[list[PlacedLine]] = []
    image_sizes: list[tuple[int, int]] = []
    for index, image in enumerate(images, start=1):
        text, placed = await recognize_page(image, engine, tessdata)
        page_txt = pages_dir / f"page_{index:04d}.txt"
        page_txt.write_text(text + "\n", encoding="utf-8")
        pages.append((index, text))
        layouts.append(placed)
        image_sizes.append(_png_size(image))
        parts.append(f"----- Page {index} -----\n{text}")
        print(f"scan {index}/{len(images)}: {len(text)} characters", flush=True)
    document = out_dir / f"{pdf.stem}.txt"
    document.write_text("\n\n".join(parts) + "\n", encoding="utf-8")
    word = out_dir / f"{pdf.stem}.docx"
    write_docx(word, pages, _pdf_page_sizes(pdf), layouts, image_sizes)
    return document, word


def main(argv: list[str] | None = None) -> int:
    args = _build_parser().parse_args(argv)
    pdf = args.pdf
    if not pdf.is_file():
        print(f"PDF not found: {pdf}", file=sys.stderr)
        return 1
    if args.dpi < 72:
        print("DPI must be at least 72.", file=sys.stderr)
        return 1
    out_dir = args.output or (Path(__file__).resolve().parent / "output" / pdf.stem)
    out_dir.mkdir(parents=True, exist_ok=True)
    document, word = asyncio.run(extract(pdf, out_dir, args.dpi))
    print(f"text: {document}")
    print(f"docx: {word}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
